"""Orquestrador: le a planilha, pede a senha e roda o pipeline por conta.

Pipeline por conta (na ordem):
  1-3. Download IMAP (.eml por UID, flags/INTERNALDATE/Message-ID em meta.json)
  4.   Construcao do .pst via Redemption (_incompleto.pst)
  5.   Verificacao por pasta e por Message-ID
  6.   Se OK: renomeia para <Nome>.pst; senao: mantem _incompleto.pst e marca FALHA
  7.   Fecha o store; limpa os .eml apenas em caso de SUCESSO (ver README).

Flags de CLI:
  --dry-run        so conecta, lista pastas e conta (nao baixa nem importa)
  --conta EMAIL    processa so uma conta
  --forcar         reprocessa uma conta ja marcada como CONCLUIDO
"""
import argparse
import configparser
import getpass
import os
import shutil
import sys

import imap_downloader
import pst_builder
import state_manager as sm
import verifier
from logger import get_logger, log_etapa, setup_logger
from progress import NullProgress, Progress
from util import sanitizar_componente


# --------------------------------------------------------------------------
# Configuracao e argumentos
# --------------------------------------------------------------------------
def _dir_base_app():
    if getattr(sys, "frozen", False):          # rodando como .exe (PyInstaller)
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def _parse_args():
    p = argparse.ArgumentParser(
        description="Backup em massa de caixas IMAP para .pst (via Redemption).")
    p.add_argument("planilha", help="Planilha (.csv ou .xlsx) com colunas: nome, email")
    p.add_argument("--config", default=None,
                   help="Caminho do config.ini (padrao: ao lado do executavel/script)")
    p.add_argument("--dry-run", action="store_true",
                   help="So conecta, lista pastas e conta (nao baixa nem importa)")
    p.add_argument("--conta", default=None, help="Processa so esta conta (email)")
    p.add_argument("--forcar", action="store_true",
                   help="Reprocessa uma conta ja marcada como CONCLUIDO")
    p.add_argument("--progress", default=None,
                   help="Arquivo JSON de progresso (usado pela interface grafica)")
    p.add_argument("--senha-stdin", action="store_true",
                   help="Le a senha da 1a linha do stdin (usado pela interface grafica)")
    return p.parse_args()


def carregar_config(caminho):
    if caminho is None:
        caminho = os.path.join(_dir_base_app(), "config.ini")
    cp = configparser.ConfigParser(interpolation=None)
    if not cp.read(caminho, encoding="utf-8"):
        raise FileNotFoundError(f"config.ini nao encontrado: {caminho}")
    return {
        "host": cp.get("imap", "host").strip(),
        "porta": cp.getint("imap", "porta", fallback=993),
        "timeout": cp.getint("imap", "timeout", fallback=60),
        "lote": cp.getint("imap", "lote", fallback=50),
        "base": cp.get("saida", "base").strip(),
        "pastas_ignorar": [s.strip() for s in
                           cp.get("pastas", "ignorar", fallback="").split(";") if s.strip()],
    }


# --------------------------------------------------------------------------
# Leitura e validacao da planilha
# --------------------------------------------------------------------------
def _norm_header(h):
    h = (h or "").strip().lower()
    if h in ("e-mail", "email", "e_mail"):
        return "email"
    if h in ("nome", "name"):
        return "nome"
    return h


def _ler_csv(caminho):
    import csv
    registros = []
    with open(caminho, "r", encoding="utf-8-sig", newline="") as f:
        amostra = f.read(2048)
        f.seek(0)
        delim = ";" if amostra.count(";") > amostra.count(",") else ","
        reader = csv.DictReader(f, delimiter=delim)
        campos = {c: _norm_header(c) for c in (reader.fieldnames or [])}
        for row in reader:
            registros.append({campos.get(k, k): v for k, v in row.items()})
    return registros


def _ler_xlsx(caminho):
    from openpyxl import load_workbook
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        ws = wb.active
        linhas = ws.iter_rows(values_only=True)
        try:
            cabecalho = next(linhas)
        except StopIteration:
            return []
        idx = {_norm_header(str(c) if c is not None else ""): i
               for i, c in enumerate(cabecalho)}
        registros = []
        for linha in linhas:
            reg = {}
            for chave in ("nome", "email"):
                if chave in idx and idx[chave] < len(linha):
                    reg[chave] = linha[idx[chave]]
            registros.append(reg)
        return registros
    finally:
        wb.close()


def ler_planilha(caminho):
    ext = os.path.splitext(caminho)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _ler_xlsx(caminho)
    return _ler_csv(caminho)


def validar_planilha(registros):
    """Remove espacos, descarta duplicados (por email) e linhas sem email,
    e atribui uma pasta de saida unica por conta (tratando homonimos)."""
    vistos = set()
    contas = []
    for linha in registros:
        nome = str(linha.get("nome") or "").strip()
        email = str(linha.get("email") or "").strip()
        if not email or "@" not in email:
            continue
        chave = email.lower()
        if chave in vistos:
            continue
        vistos.add(chave)
        if not nome:
            nome = email.split("@")[0]
        contas.append({"nome": nome, "email": email})

    # Homonimos -> usar "Nome (email)"; garante unicidade final com sufixo.
    base_por_conta = [sanitizar_componente(c["nome"]) for c in contas]
    repetido = {}
    for b in base_por_conta:
        repetido[b.lower()] = repetido.get(b.lower(), 0) + 1

    usados = {}
    for c, base in zip(contas, base_por_conta):
        if repetido[base.lower()] > 1:
            pasta = sanitizar_componente(f"{c['nome']} ({c['email']})")
        else:
            pasta = base
        chave = pasta.lower()
        if chave in usados:
            usados[chave] += 1
            pasta = f"{pasta} ({usados[chave]})"
        else:
            usados[chave] = 1
        c["pasta"] = pasta
    return contas


# --------------------------------------------------------------------------
# Limpeza
# --------------------------------------------------------------------------
def _limpar_temp(dir_eml):
    if os.path.isdir(dir_eml):
        shutil.rmtree(dir_eml, ignore_errors=True)


def _limpar_tudo(dir_conta, dir_eml, pst_incompleto, pst_final):
    _limpar_temp(dir_eml)
    for p in (pst_incompleto, pst_final):
        try:
            pst_builder.remover_com_retentativa(p)
        except Exception:
            pass
    for p in (os.path.join(dir_conta, "state.json"),
              os.path.join(dir_conta, "meta.json")):
        try:
            if os.path.exists(p):
                os.remove(p)
        except OSError:
            pass


# --------------------------------------------------------------------------
# Pipeline por conta
# --------------------------------------------------------------------------
def processar_conta(cfg, reg, senha, dry_run, forcar, sessao, progress=None):
    progress = progress or NullProgress()
    email = reg["email"]
    nome = reg["nome"]
    dir_conta = os.path.join(cfg["base"], reg["pasta"])
    dir_eml = os.path.join(dir_conta, "_temp")
    pst_incompleto = os.path.join(dir_conta, "_incompleto.pst")
    pst_final = os.path.join(dir_conta, f"{sanitizar_componente(nome)}.pst")

    estado = sm.carregar_estado(dir_conta)
    estado["email"] = email
    estado["nome"] = nome
    meta = sm.carregar_meta(dir_conta)

    if estado["status"] == sm.CONCLUIDO and not forcar and not dry_run:
        log_etapa(email, "pipeline", "PULADA", "ja concluida (use --forcar para refazer)")
        return "PULADA"

    if forcar and not dry_run:
        _limpar_tudo(dir_conta, dir_eml, pst_incompleto, pst_final)
        estado = sm.carregar_estado(dir_conta)
        estado["email"] = email
        estado["nome"] = nome
        meta = {}

    if not dry_run:
        os.makedirs(dir_conta, exist_ok=True)

    # ---- 1-3. DOWNLOAD ----
    if not dry_run:
        sm.marcar_status(dir_conta, estado, sm.BAIXANDO)
    log_etapa(email, "login_imap", "INICIO", f"host={cfg['host']} porta={cfg['porta']}")

    def salvar():
        sm.salvar_estado(dir_conta, estado)
        sm.salvar_meta(dir_conta, meta)

    resumo = imap_downloader.executar_download(
        cfg, email, senha, dir_conta, dir_eml, estado, meta,
        dry_run, salvar=(None if dry_run else salvar), progress=progress,
    )
    log_etapa(email, "download", "RESUMO",
              f"pastas={len(resumo)} mensagens_servidor={sum(resumo.values())}")

    if dry_run:
        log_etapa(email, "pipeline", "DRY_RUN_OK",
                  "apenas listagem/contagem; nada baixado nem importado")
        return "DRY_RUN"

    salvar()
    sm.marcar_status(dir_conta, estado, sm.BAIXADO)

    # ---- 4-5. CONSTRUIR PST + VERIFICAR ----
    builder = pst_builder.PstBuilder(sessao, email, nome)
    try:
        sm.marcar_status(dir_conta, estado, sm.CONSTRUINDO_PST)
        # Rebuild limpo: um _incompleto.pst de uma tentativa anterior e descartado
        # (os .eml completos permitem reconstruir sem duplicar mensagens).
        pst_builder.remover_com_retentativa(pst_incompleto)
        builder.abrir_store(pst_incompleto)
        progress.start_build(sum(len(ep.get("uids_baixados", []))
                                 for ep in estado["pastas"].values()))
        importados = builder.importar(dir_eml, estado, meta, progress=progress)
        log_etapa(email, "importar", "OK", f"importados={importados}")

        sm.marcar_status(dir_conta, estado, sm.VERIFICANDO)
        progress.set_phase("verificando")
        indice = builder.indexar()
        ok, relatorio = verifier.verificar(email, estado, meta, indice)
    finally:
        builder.fechar()

    # ---- 6-7. RENOMEAR / LIMPAR ----
    if ok:
        pst_builder.renomear_com_retentativa(pst_incompleto, pst_final)
        sm.marcar_status(dir_conta, estado, sm.CONCLUIDO)
        log_etapa(email, "pipeline", "CONCLUIDO", f"pst={os.path.basename(pst_final)}")
        _limpar_temp(dir_eml)   # limpeza dos .eml so no sucesso
        log_etapa(email, "limpeza_temp", "OK")
        return "CONCLUIDO"

    sm.marcar_status(dir_conta, estado, sm.FALHA)
    log_etapa(email, "pipeline", "FALHA",
              f"verificacao divergente; mantido {os.path.basename(pst_incompleto)}; "
              f"faltantes={relatorio['faltantes_total']}")
    return "FALHA"


# --------------------------------------------------------------------------
def main():
    args = _parse_args()
    cfg = carregar_config(args.config)
    setup_logger(os.path.join(cfg["base"], "logs"))
    logger = get_logger()

    if not os.path.exists(args.planilha):
        logger.error("Planilha nao encontrada: %s", args.planilha)
        return 2

    contas = validar_planilha(ler_planilha(args.planilha))
    if args.conta:
        contas = [c for c in contas if c["email"].lower() == args.conta.lower()]
        if not contas:
            logger.error("Conta %s nao encontrada na planilha.", args.conta)
            return 2

    logger.info("Contas a processar: %d | dry_run=%s | forcar=%s | base=%s",
                len(contas), args.dry_run, args.forcar, cfg["base"])

    if args.senha_stdin:
        senha = sys.stdin.readline().rstrip("\r\n")
    else:
        senha = getpass.getpass("Senha comum das caixas (nao sera exibida nem gravada): ")
    if not senha:
        logger.error("Senha vazia; abortando.")
        return 2

    progress = Progress(args.progress, len(contas)) if args.progress else NullProgress()
    sessao = pst_builder.RedemptionSession()  # UMA sessao MAPI para todo o processo
    resultados = {}
    try:
        for idx, c in enumerate(contas, start=1):
            try:
                progress.start_account(idx, c["email"])
                resultados[c["email"]] = processar_conta(
                    cfg, c, senha, args.dry_run, args.forcar, sessao, progress)
                progress.finish_account(resultados[c["email"]])
            except KeyboardInterrupt:
                logger.warning(
                    "Interrompido (Ctrl+C). Estado salvo; reexecute para retomar.")
                resultados[c["email"]] = "INTERROMPIDO"
                progress.finish_account("INTERROMPIDO")
                raise
            except Exception as e:
                resultados[c["email"]] = "ERRO"
                progress.finish_account("ERRO")
                log_etapa(c["email"], "pipeline", "ERRO", f"{type(e).__name__}: {e}")
                logger.exception("Erro ao processar %s", c["email"])
                # Isolamento por conta: segue para a proxima.
    finally:
        senha = None  # best-effort: nao manter a senha em memoria apos o uso
        sessao.encerrar()  # Logoff + CoUninitialize UMA vez, no fim
        progress.finish_all()
        logger.info("===== RESUMO FINAL =====")
        for email, r in resultados.items():
            logger.info("  %-40s -> %s", email, r)

    return 0


if __name__ == "__main__":
    sys.exit(main())
